"""Beginner team handbook, created using the bundled document runtime."""
from pathlib import Path
import math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs'/'Shirt_Recommendation_Project_Team_Guide.docx'
QA=ROOT.parent/'JUNK created by qoder'/'team_guide_qa'
QA.mkdir(parents=True,exist_ok=True)

def architecture():
    im=Image.new('RGB',(1800,1330),'white');d=ImageDraw.Draw(im)
    font=lambda n,b=False:ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf' if b else 'C:/Windows/Fonts/arial.ttf',n)
    black='#17211c';green='#416653';line='#6a7b70'
    def centered(x,y,text,size=30,b=False):
        for row in text.split('\n'):
            box=d.textbbox((0,0),row,font=font(size,b))
            d.text((x-(box[2]-box[0])/2,y),row,font=font(size,b),fill=black)
            y+=size+9
    def box(rect,title,body,role):
        x0,y0,x1,y1=rect
        d.rounded_rectangle(rect,radius=18,fill='#f1f5f1',outline=green,width=3)
        centered((x0+x1)/2,y0+22,title,29,True)
        centered((x0+x1)/2,y0+78,body,26)
        centered((x0+x1)/2,y1-46,role,22)
    def arrow(points,dashed=False):
        for a,b in zip(points,points[1:]):
            if dashed:
                length=math.dist(a,b)
                for t in range(0,int(length),22):
                    f=t/length;g=min(1,(t+11)/length)
                    d.line((a[0]+f*(b[0]-a[0]),a[1]+f*(b[1]-a[1]),a[0]+g*(b[0]-a[0]),a[1]+g*(b[1]-a[1])),fill=line,width=4)
            else:d.line([a,b],fill=line,width=5)
        a,b=points[-2:];angle=math.atan2(b[1]-a[1],b[0]-a[0]);length=17
        wings=[(b[0]-length*math.cos(angle+s*.55),b[1]-length*math.sin(angle+s*.55)) for s in [-1,1]]
        d.polygon([b,*wings],fill=line)
    d.text((55,25),'Prepare the catalog once',font=font(36,True),fill=black)
    specs=[('Amazon metadata','Product details\nand image URLs','Member 1'),
           ('Shirt selection','CLIP screening\nand data cleaning','Member 1'),
           ('Kaggle vision','Shirt masks\nColors and patterns','Member 2'),
           ('Merge by ID','Match parent_asin\nKeep status labels','Members 1 and 4'),
           ('Master Parquet','50,224 products\nReady for search','Member 4')]
    for i,(title,body,role) in enumerate(specs):
        x=55+i*342;box((x,110,x+295,325),title,body,role)
        if i<4:arrow([(x+295,217),(x+342,217)])
    d.text((55,405),'Search the prepared catalog for each user request',font=font(36,True),fill=black)
    boxes=[((55,505,555,730),'User query','Example plain black shirt\nunder 1500 rupees','Member 5'),
           ((655,505,1155,730),'FastAPI backend','Send query and schema\nReceive a SQL plan','Members 3 and 5'),
           ((1255,505,1755,730),'NVIDIA hosted LLM','Nemotron 3.5 Lightning\nGenerate SQL as JSON','Member 3'),
           ((1255,935,1755,1160),'SQL validation','Read queries only\nApply exact color rules','Members 3 and 4'),
           ((655,935,1155,1160),'DuckDB on Parquet','Filter matching product IDs\nKeep all requested conditions','Member 4'),
           ((55,935,555,1160),'Rank and display','Sort matching products\nShow images and status','Members 4 and 5')]
    for args in boxes:box(*args)
    for points in [[(555,617),(655,617)],[(1155,617),(1255,617)],[(1505,730),(1505,935)],
                   [(1255,1047),(1155,1047)],[(655,1047),(555,1047)]]:arrow(points)
    arrow([(1570,325),(1780,325),(1780,1260),(905,1260),(905,1160)],True)
    d.text((1050,1280),'Dashed arrow is the saved catalog file',font=font(25),fill=black)
    p=QA/'architecture.png';im.save(p);return p

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
sec.footer_distance=Inches(.3)
for name in ['Normal','Body Text','List Bullet','List Number']:
    style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(11)
    style.font.color.rgb=RGBColor(0,0,0);style.paragraph_format.line_spacing=1.1
    style.paragraph_format.space_after=Pt(6)
for name,size in [('Title',25),('Subtitle',12),('Heading 1',19),('Heading 2',12.5)]:
    s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_before=Pt(10 if name=='Heading 2' else 0)
    s.paragraph_format.space_after=Pt(7);s.paragraph_format.keep_with_next=True
for name in ['Heading 1','Heading 2']:doc.styles[name].font.bold=True
# Remove the stock Word template's blue title rule and theme overrides.
for border in list(doc.styles.element.xpath('.//w:pBdr')):border.getparent().remove(border)
for name in ['Title','Subtitle','Heading 1','Heading 2']:
    for color in doc.styles[name].element.xpath('.//w:color'):
        for attr in ['themeColor','themeTint','themeShade']:color.attrib.pop(qn('w:'+attr),None)
        color.set(qn('w:val'),'000000')
doc.styles['Title'].font.underline=False
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Team study guide   |   Page ');r.font.name='Arial';r.font.size=Pt(9)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
doc.core_properties.title='Shirt Recommendation Project Team Guide'
doc.core_properties.subject='Five team roles architecture and technology stack'
doc.core_properties.author='Project team'

def p(text,style=None):return doc.add_paragraph(text,style)
def h(text,level=2):return doc.add_heading(text,level)
def lead(label,text):
    para=doc.add_paragraph();para.add_run(label+' ').bold=True;para.add_run(text);return para
def bullet(text):return p(text,'List Bullet')
def qa(question,answer):lead(question,answer)
def page(title):h(title,1).paragraph_format.page_break_before=True
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,value in enumerate(headers):t.rows[0].cells[i].text=value
    for row in rows:
        cells=t.add_row().cells
        for i,value in enumerate(row):cells[i].text=str(value)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side,val in [('top','95'),('bottom','95'),('left','115'),('right','115')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');m.append(e)
            tcpr.append(m)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EEF0' if ri==0 else ('F8FAFA' if ri%2==0 else 'FFFFFF'));tcpr.append(sh)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.05
                for run in para.runs:
                    run.font.size=Pt(10.5);run.font.color.rgb=RGBColor(0,0,0);run.bold=(ri==0)
    p('');return t

p('Shirt Recommendation Project Team Guide','Title')
p('Simple explanations for five members preparing a presentation and judge discussion','Subtitle')
p('Our project lets a user describe a shirt in normal language. A hosted language model turns that request into SQL. DuckDB searches our prepared shirt catalog, and the website shows ranked matching products with images.')
p('This guide explains what each member should understand, the work belonging to each role, and simple answers for a judge. Replace the member numbers with your names. Use the sample scripts for the work you actually handled; shared work can be explained together.')
h('The project in one example')
lead('User request','Plain black shirt under 1500 rupees.')
lead('Search conditions','Color is only black, pattern is plain, and the catalog price is below INR 1,500. Shirts tagged navy|black or burgundy|black are excluded.')
lead('Output','Matching shirts shown as cards containing product ID, title, image, price, rating and attribute status.')
h('Five member responsibility map')
table(['Member','Main responsibility','What they hand to the next stage'],[
('1','Dataset preparation','Clean shirt catalog with product IDs and image URLs'),
('2','Image attribute extraction','Colors patterns and uncertainty status per product'),
('3','Language model and SQL','A SQL filter based on the user request'),
('4','DuckDB and ranking','Matching products in a useful order'),
('5','Application and testing','A working search interface and tested demonstration')],[.9,2.4,3.7])
h('Numbers everyone should know')
p('50,224 products share 31,793 unique image URLs. After the merge, 49,658 products have nonempty color tags and 41,752 have known patterns. These are coverage counts, not accuracy percentages. Some predictions remain uncertain or unknown.')
lead('Project folder',r'C:\Users\thopa\OneDrive\Desktop\project\REAL SYSTEM\DuckDB idea')

page('Project architecture')
p('The top part prepares the data before searching. The lower part runs whenever a user submits a request. Product images are not reprocessed by the vision models for every search.')
pic=architecture();para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=para.add_run();run.add_picture(str(pic),width=Inches(7))
inline=run._r.xpath('.//wp:docPr')[0];inline.set('descr','Two stage project architecture showing catalog preparation and the live query flow with each team role')
h('How to explain the diagram')
bullet('First we create a shirt catalog, run image models on Kaggle, and merge their results using parent_asin.')
bullet('During search, NVIDIA generates SQL. Our backend checks it before DuckDB reads the Parquet file.')
bullet('DuckDB finds products that meet the filters. The app ranks those products and shows their images and details.')
lead('Important distinction','The vision models identify shirt attributes before search. The live NVIDIA language model interprets the request; it does not look at product images in this query flow.')

roles=[
('Role 1 Dataset preparation',
 'Prepare a clean list of shirts so every later stage works on the same products.',
 ["Understand the Amazon 2023 product metadata: titles, brands, ratings and image links.",
  'Understand how the earlier CLIP screening selected shirts from the larger clothing catalog.',
  'Keep one record per parent_asin, check duplicate or missing IDs, and retain existing price, rating and review-derived fields.',
  'Prepare the master Parquet and its CSV copy. Keep IDs and image URLs consistent during the vision merge.'],
 [('Metadata','Information about a product, such as its title, brand and image URL.'),
  ('parent_asin','The parent product identifier. We use it like a roll number to match records across files.'),
  ('Parquet','A compact table file that DuckDB can query. CSV is a plain table copy that is easy to inspect.')],
 'My responsibility is dataset preparation. I work with the Amazon product metadata and the shirt-selection output. I check product IDs, organize useful fields and prepare the master catalog. The final catalog contains 50,224 products. Its stable product IDs let us attach image-based attributes without mixing up products.',
 [('Why do you need parent_asin?','It connects a product to its color and pattern results. Row positions can change, but an ID match keeps the right attributes with the right product.'),
  ('Did you collect live Amazon products?','No. This is a research catalog derived from the Amazon 2023 dataset. The images are accessed through the saved URLs.')],
 'data/shirts_master.parquet, data/exports/shirts_master.csv, merge_vision.py and data/reports/vision_merge_summary.json.',
 'Product titles and old color tags are not always correct. A parent product can also have variants; the selected image does not describe every possible variant.'),
('Role 2 Image based color and pattern extraction',
 'Use the product image to estimate the colors and pattern of the shirt fabric.',
 ['Run the v3.2 image pipeline on Kaggle using two T4 GPUs. This uses pretrained models for prediction, not new model training.',
  'Locate the shirt with Grounding DINO and get its mask with SAM. SegFormer supplies supporting clothing information.',
  'Name colors from retained shirt pixels using a color palette. Use CLIP plus a stripe-direction check for the pattern.',
  'Save predictions, product IDs, model signals and failure or uncertainty status. Process each identical image URL once.'],
 [('Mask','A selection of image pixels believed to belong to the shirt. It is meant to exclude the background and other objects.'),
  ('Inference','Using a model that is already trained to make a prediction.'),
  ('needs_review','The script found uncertainty. It can still contain predictions, but it is not a verified label.')],
 'My responsibility is image attribute extraction. I run the models on Kaggle to isolate shirts and estimate their colors and patterns. The script reuses results for shared images, so 31,793 unique images cover 50,224 products. We retain uncertainty labels and avoid inventing colors when the mask fails.',
 [('Why not use the whole image for color?','It may contain pants, skin, a tie or a background. Those colors should not be treated as shirt fabric colors.'),
  ('Are gold and yellow merged?','No. The palette keeps them separate, just as it keeps maroon separate from red. The prediction can still be wrong because of lighting or masking.')],
 'Kaggle v3.2 vision_pipeline.py, attribute_core.py, shirt_attributes.py and the final shirt_results_backup ZIP.',
 'Masks can include accessories. Shadows can add extra shade names. The normal color list uses an 8 percent pixel-share threshold, so small printed colors may be omitted; raw candidates are retained.'),
('Role 3 Language model and SQL generation',
 'Turn the user\'s ordinary sentence into a database search instruction.',
 ['Connect to NVIDIA\'s hosted Nemotron 3.5 Lightning model through its API.',
  'Send the request along with available columns, allowed attribute names and examples. The entire 50k catalog is not sent to the model.',
  'Ask the model for JSON containing SQL, an explanation, unsupported details and a sorting preference.',
  'Work with Member 4 to validate generated SQL and normalize color and pattern wording before execution.'],
 [('LLM','A large language model. Here it interprets a sentence and writes SQL.'),
  ('SQL','A language for selecting records from a table using conditions.'),
  ('API','The connection through which our Python backend sends a request to NVIDIA and receives the model response.')],
 'My responsibility is the natural-language query layer. I give the NVIDIA model the user request and our catalog schema. It generates a structured SQL plan. We check this plan and use the actual dataset vocabulary, so words such as striped map to the three stored stripe patterns. The model does not invent product records.',
 [('Does the model choose products from memory?','No. It writes the filter. DuckDB selects actual records from our own Parquet catalog.'),
  ('What happens with black shirts with white buttons?','There is no button_color field. The app reports that detail as unsupported instead of claiming that it applied the condition.')],
 'llm_sql.py, app.py and tests/test_query_app.py. The private .env contains API settings and must not appear in slides.',
 'The model can generate incorrect SQL or misunderstand a request. SQL checks reduce execution and vocabulary errors, but they do not prove that every request is interpreted perfectly.'),
('Role 4 DuckDB filtering and recommendation ranking',
 'Execute valid filters on the prepared data and put matching products in a useful order.',
 ['Merge vision results with the master using an exact parent_asin and image-URL match. Keep a backup of the original file.',
  'Create a DuckDB view of the Parquet catalog. Restrict queries to the shirt table and the configured data file.',
  'Apply color, pattern, sleeve and budget filters. For plain black, enforce only the black tag; mixed tags are excluded.',
  'Count all matches, rank them and return the requested number of product cards. Preserve historical or estimated price labels.'],
 [('DuckDB','The SQL engine that reads and searches our table file. A separate database server is not needed.'),
  ('Filtering','Removing products that do not satisfy the requested conditions.'),
  ('Ranking','Sorting the remaining products. It does not bring back products that failed the filters.')],
 'My responsibility is the data query and ranking layer. DuckDB reads the updated Parquet file and applies the validated conditions. For plain black under 1500, I require only black, plain pattern and the price limit. I rank only matching products. The current default favors historical exact-price records, then a rating score that accounts for rating count.',
 [('Why use DuckDB?','It can run SQL directly on Parquet and is simple to use for this local catalog. We do not have to copy the data into a separate database server.'),
  ('How are results ranked?','The default uses price-source tiers and a rating average adjusted with a 20-review prior. Requests for cheapest, highest-rated or most reviewed use separate sorts.')],
 'query_engine.py, request_policy.py, merge_vision.py and data/shirts_master.parquet.',
 'Strict single-color filtering may exclude a truly black shirt if its prediction includes a shadow shade. Prices are historical or estimated, not live Amazon prices.'),
('Role 5 Application integration and testing',
 'Give users a simple search page and check that the complete system works together.',
 ['Build the HTML, CSS and JavaScript search page and connect it to FastAPI routes.',
  'Show product images, titles, IDs, prices, ratings, color tags, patterns and uncertainty badges.',
  'Use Uvicorn to run the local app. Handle API errors, missing images, unsupported details and empty results clearly.',
  'Test merge alignment, SQL restrictions, exact colors, API behavior and real searches. Prepare a repeatable demonstration.'],
 [('Frontend','The page the user sees and clicks.'),
  ('Backend','The Python code that receives the request, calls the model and searches the data.'),
  ('Test','A check with an expected outcome, such as proving that a plain black result has no other color tag.')],
 'My responsibility is the interface and end-to-end integration. I connect the search form to the backend and show returned products as cards. I display estimated prices and uncertain attributes clearly. I test that valid requests work and unsafe SQL is rejected. The latest test run passed 18 automated checks, and live requests were also tested.',
 [('How do you test the whole flow?','I send a real request through the API and check the returned records against the requested colors, pattern and budget. Automated tests also cover failures and invalid SQL.'),
  ('What if there are no matches?','The page shows no matching products. It does not silently remove a budget or color condition to fill the screen.')],
 'app.py, web/index.html, launch.py, start.bat and the tests folder.',
 'Passing software tests does not establish the visual accuracy of every color or pattern tag. Live stock, delivery and personalized recommendations are not implemented.')]

for title,job,tasks,terms,script,questions,files,limitation in roles:
    page(title)
    lead('Assigned member','________________________________')
    h('Your job');p(job)
    h('Work you should understand')
    for task in tasks:bullet(task)
    h('Words in simple language')
    for term,meaning in terms:lead(term,meaning)
    h('Short presentation script');p(script)
    h('Questions the judge may ask')
    for question,answer in questions:qa(question,answer)
    lead('Files to show',files)
    lead('Limitation to explain',limitation)

page('Demonstration and shared judge questions')
h('A simple demonstration')
for text in ['Start the app with start.bat and open http://127.0.0.1:8010.',
 'Search for plain black shirt under 1500. In the latest checked request, 20 products matched; all returned records were tagged only black, plain and below INR 1,500.',
 'Show a card and explain its image, product ID, price source and uncertainty status. Open the SQL details to explain which filters were used.',
 'Try blue and white striped shirts under 2000 rupees to show a multi-color request. One earlier live check found 10 matching records.',
 'Try black shirts with white buttons to demonstrate an unsupported field. Finish by explaining remaining limitations.']:
    p(text,'List Number')
h('Answers all five members should know')
qa('Did you train all the models?','The vision models and NVIDIA language model are pretrained. The earlier price-preparation stage trained a separate regression model on known prices. The search app uses its saved price estimates.')
qa('Are all 50k products tagged perfectly?','No. Some tags are uncertain and some fields remain unknown. Coverage and successful processing are not accuracy measurements.')
qa('Is this the same system as the research paper?','The paper is a reference for apparel recommendation. Our current implementation uses an LLM-to-SQL and DuckDB approach. It does not reproduce the paper\'s full cascaded review-embedding pipeline.')
qa('Is this Amazon or Flipkart\'s production architecture?','No. It is our research prototype over an Amazon-derived catalog. We should not claim knowledge of their private production systems.')
qa('How do you prevent unwanted SQL operations?','The backend permits a restricted read query, checks table and column names, blocks file readers and modifications, and restricts DuckDB\'s file access.')
qa('What is still missing?','Live inventory, current prices, button color, exact size availability, personalization and a labeled evaluation of recommendation quality.')
qa('What does a price estimate mean?','Many prices are estimated from the existing dataset-building stage. Cards label estimates, and INR values use the project\'s 83 INR per USD conversion.')
lead('How to rehearse','Each member should explain their input, main steps, output and one limitation. Then practice handing the next stage to the next member. Understand the diagram rather than memorizing model names alone.')

page('Technology stack for the live application')
p('These technologies run the current query application. They are separate from the image models used earlier to prepare the catalog.')
table(['Technology','What it does','Where we use it'],[
('Python','Main programming language','Backend, merging, data checks and tests'),
('NVIDIA hosted API','Runs the language model remotely','Interprets the request and produces the SQL plan'),
('Nemotron 3.5 Lightning','The current text language model','Model ID nvidia/nemotron-3.5-lightning-30b-a3b; no local GPU is needed for the API call'),
('SQL','Describes search conditions','Select matching product IDs using catalog fields'),
('SQLGlot','Parses SQL into a structure we can check','Rejects unsupported operations and helps normalize attribute filters'),
('DuckDB','Runs SQL on the Parquet file','Filters the catalog and supports deterministic ranking'),
('Apache Parquet and PyArrow','Store and read the catalog efficiently','shirts_master.parquet and data exports'),
('Pandas','Works with tables in Python','Merge attributes by ID, check data and prepare results'),
('FastAPI and Pydantic','Provide backend routes and check request structure','Search, browsing and health endpoints'),
('Uvicorn','Runs the local FastAPI server','Serves the app at port 8010'),
('HTML CSS JavaScript','Build the page and handle interactions','Search form, loading states and product cards'),
('Requests and JSON','Send API calls and exchange structured data','Python to NVIDIA and browser to backend'),
('python dotenv','Loads private local settings','Provider URL, model and API key from .env'),
('unittest and TestClient','Automated software checks','Merge, filter, SQL, API and integration tests')],[1.6,2.3,3.1])
lead('API access','NVIDIA provides hosted access for prototyping, with account and model rate limits. A successful test does not establish unlimited free production access.')
lead('Credential handling','The API key is stored in the local ignored .env file. Never put its value in a slide, screenshot, shared code file or document.')

page('Technology stack for preparing the catalog')
table(['Technology or model','Simple purpose','Role in this project'],[
('Amazon 2023 metadata','Source product information','Starting data for the curated shirt catalog'),
('CLIP ViT B 32','Compares images with shirt descriptions','Earlier shirt-selection screening from the larger clothing dataset'),
('Kaggle notebooks and T4 x2','Provide cloud GPUs','Run the v3.2 vision script over the unique product images'),
('Grounding DINO base','Finds the shirt location','Detects a shirt bounding box in a product image'),
('SAM ViT base','Selects object pixels','Creates the main box-prompted shirt mask in v3.2'),
('SegFormer B2 clothes','Recognizes clothing regions','Supporting overlap and clothing-region information; it is not the sole folded-shirt mask'),
('CLIP ViT B 16','Compares shirt images with pattern descriptions','Predicts the shirt pattern from the masked image'),
('CIELAB and CIEDE2000','Compare photographic colors','Name histogram colors using a 38-name shade palette'),
('Image gradient analysis','Measures visible line direction','Supports horizontal vertical or diagonal stripe decisions'),
('PyTorch and Transformers','Run pretrained image models','GPU inference and image-model processing'),
('Hugging Face Hub','Provides model files and revisions','Downloads and records model snapshots'),
('NumPy SciPy Pillow','Process pixels arrays and images','Masks, color statistics, stripe checks and previews'),
('SQLite','Stores progress locally','Per-image checkpointing and resume without repeating completed URLs'),
('CSV ZIP and JSON','Make outputs portable and inspectable','Results, backups, settings, vocabularies and summaries')],[1.7,2.2,3.1])
lead('Training versus prediction','The Kaggle run used existing model weights. Our 50k catalog was processed for predictions; those models were not trained from scratch on it.')
lead('Earlier experiments','Qwen2.5 VL and Keypoint R CNN were tried in the slower experimental pipeline. They are not active components of the final v3.2 catalog run or the current live search path.')
lead('Current boundaries','The app uses existing review and rating fields, but it does not run a live sentiment model or the legacy cascaded review-embedding system for each request.')

page('Technology stack for earlier dataset enrichment')
p('The master catalog already contained price and text-derived fields before the final color and pattern merge. These earlier tools help explain where those fields came from. They are not rerun for every search.')
table(['Technology','Simple purpose','Existing use in the catalog'],[
('scikit learn','Tools for conventional machine learning','Preparing and training the earlier price-estimation pipeline'),
('Gradient boosting','Learns how input details relate to a number','HistGradientBoostingRegressor estimates missing USD prices from product features'),
('TF IDF','Turns words into numeric features','Represents product-title text for the earlier price model'),
('Encoders and imputers','Prepare categories and missing values','Process brand and other features before prediction'),
('Joblib','Saves and reloads a Python model','Stores the price-model bundle in data/models/price_model.joblib'),
('Regular expressions','Find words and phrases in text','Extract fabric, fit and pocket clues; mine fit feedback from reviews'),
('JSONL and gzip','Store large records in compressed files','Original Amazon metadata and review dataset formats')],[1.6,2.3,3.1])
h('Explain price estimation simply')
p('Some products have a price in the original metadata. For missing prices, the earlier estimator learns from products with known prices and predicts a value using features such as title and brand. An estimated value remains labeled estimated; it does not become an actual current Amazon price.')
h('Explain review features simply')
p('Earlier scripts search product text and review phrases for clues such as cotton, regular fit, runs small or true to size. These saved attributes are available in the catalog. This is different from running a new review-language model during every search.')
lead('Relevant files','scripts/legacy/03b_price_estimator_v2.py, scripts/legacy/04_job_a_attributes.py, scripts/legacy/06_job_c_review_mining.py and data/models/price_model.joblib.')
lead('Keep the distinction clear','Kaggle vision predicts shirt colors and patterns. The NVIDIA language model writes SQL. The earlier regression model estimates missing prices. Each model has a different job.')

doc.save(OUT)
print(OUT)
print('Pages planned: 11; roles: 5; architecture image embedded; credential values excluded.')
